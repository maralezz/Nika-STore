"""
🔵 NIKA Сантехника — интернет-магазин
"""
import os
import io
import csv
import json
import re
import shutil
import sqlite3
import time
import zipfile
import threading
import smtplib
from concurrent.futures import ThreadPoolExecutor, as_completed
from functools import wraps
from urllib.parse import urlparse, unquote, urlencode
from xml.etree import ElementTree as ET
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart

import requests
from bs4 import BeautifulSoup
from flask import (Flask, render_template, request, redirect, url_for,
                   session, flash, g, jsonify, send_file)
from werkzeug.security import generate_password_hash, check_password_hash
from openpyxl import load_workbook
import xlrd

from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from PIL import Image

app = Flask(__name__)
app.secret_key = 'nika-aquastyle-secret-key-2026'

def product_display_name(p):
    try: name = p['name'] or ''
    except: return ''
    try: sku = p['sku'] or ''
    except: sku = ''
    if not sku or not name: return name

    def cleanup(text):
        text = re.sub(r'\s+', ' ', text).strip()
        text = re.sub(r'\s+([,.])', r'\1', text)
        text = re.sub(r'([,.])\s+([,.])', r'\1', text)
        text = re.sub(r'^\s*[,.\-]+\s*', '', text)
        text = re.sub(r'\s*[,.\-]+\s*$', '', text)
        return text.strip()

    s = str(sku).strip()
    if len(s) < 3: return name

    pattern = r'(?<![A-Za-zА-Яа-я0-9])' + re.escape(s) + r'(?![A-Za-zА-Яа-я0-9])'
    new_name = re.sub(pattern, '', name, flags=re.IGNORECASE)
    if new_name != name:
        cleaned = cleanup(new_name)
        return cleaned if cleaned else name

    tokens = [t for t in re.split(r'[-_/\s]+', s) if len(t) >= 4]
    tokens.sort(key=len, reverse=True)

    for token in tokens:
        pattern_parts = []
        for i, ch in enumerate(token):
            pattern_parts.append(re.escape(ch))
            if i < len(token) - 1:
                nxt = token[i + 1]
                if ch.isalpha() != nxt.isalpha(): pattern_parts.append(r'[\s\-]?')
        flex_pattern = (r'(?<![A-Za-zА-Яа-я0-9])' + ''.join(pattern_parts) + r'(?![A-Za-zА-Яа-я0-9])')
        new_name = re.sub(flex_pattern, '', name, flags=re.IGNORECASE)
        if new_name != name:
            cleaned = cleanup(new_name)
            return cleaned if cleaned else name
    return name

app.jinja_env.globals['product_display_name'] = product_display_name

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATABASE = os.path.join(BASE_DIR, 'nika.db')
IMG_DIR = os.path.join(BASE_DIR, 'static', 'product_images')
IMG_URL_PREFIX = '/static/product_images/'

HEADERS = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36'}
_REQUEST_DELAY = 1.6
_last_request_ts = {}
_DEBUG_COUNT = [0]
_DEBUG_LIMIT = 3

_NON_IMAGE_EXT = ('.pdf', '.dwg', '.zip', '.rar', '.7z', '.tar', '.gz', '.doc', '.docx', '.xls', '.xlsx', '.ppt', '.pptx', '.txt', '.csv', '.xml', '.json', '.html', '.mp4', '.mov', '.webm', '.avi', '.mkv', '.obj', '.3ds', '.fbx', '.max', '.stl', '.blend')
_IMG_EXT = ('.jpg', '.jpeg', '.png', '.webp', '.gif', '.bmp', '.svg', '.avif')
_DOC_EXT = ('.pdf', '.dwg', '.doc', '.docx', '.xls', '.xlsx', '.ppt', '.pptx', '.zip', '.obj', '.3ds', '.fbx', '.max', '.stl', '.blend')

_SCHEMA_URL_MARKERS = ('схем', 'черт', 'чертёж', 'drawing', 'schema', 'scheme', 'pdw', 'ptdw', 'габарит', 'размер', 'dimension', 'размеры')
_YANDEX_DISK_HOSTS = ('disk.yandex.ru', 'yadi.sk', 'disk.yandex.com')
_YANDEX_TIMEOUT = 5
_YANDEX_WORKERS = 12
_YANDEX_CACHE = {}
_YANDEX_CACHE_LOCK = threading.Lock()

_FIX_STATE = {'running': False, 'total': 0, 'processed': 0, 'resolved': 0, 'removed': 0, 'yandex_total': 0, 'yandex_done': 0, 'yandex_resolved': 0, 'yandex_failed': 0, 'started_at': None, 'finished_at': None, 'stage': 'idle', 'message': '', 'error': None}
_FIX_LOCK = threading.Lock()
_WHITE_IMAGE_CACHE = {}

def _is_yandex_disk_url(u): return u and isinstance(u, str) and any(h in u.lower() for h in _YANDEX_DISK_HOSTS)

def _url_ext(u):
    if not u or not isinstance(u, str): return ''
    low = u.lower().split('?')[0].split('#')[0]
    for e in _IMG_EXT + _DOC_EXT:
        if low.endswith(e): return e
    return ''

def _url_is_image_by_ext(u): return _url_ext(u) in _IMG_EXT
def _url_is_doc_by_ext(u): return _url_ext(u) in _DOC_EXT

def _url_looks_like_schema(u): return u and isinstance(u, str) and any(m in u.lower() for m in _SCHEMA_URL_MARKERS)

def _split_into_urls(s):
    if not s or not isinstance(s, str): return []
    text = re.sub(r'[;\t]+[ \t]*', '\n', re.sub(r',\s+', '\n', s))
    urls = []
    for line in text.split('\n'):
        line = line.strip().rstrip(';,.').rstrip(')')
        if not line: continue
        m = re.search(r'https?://', line)
        if not m: continue
        url = line[m.start():].strip()
        idx = url[8:].lower().find('http')
        if idx != -1: url = url[:idx + 8].strip()
        if url: urls.append(url)
    return urls

def _is_composite_url(u):
    if not u or not isinstance(u, str): return True
    if '\n' in u or '\r' in u or '\t' in u: return True
    if u.lower().count('http://') + u.lower().count('https://') > 1: return True
    if not u.startswith('http'): return True
    return False

def _explode_composite(u):
    if not u or not isinstance(u, str): return []
    if not _is_composite_url(u): return [u]
    return _split_into_urls(u)

def _is_valid_image_url(u):
    if not u or not isinstance(u, str): return False
    if not u.startswith('http') and not u.startswith('/static/'): return False
    if _is_yandex_disk_url(u) or _is_composite_url(u): return False
    if u.lower().split('?')[0].split('#')[0].endswith(_NON_IMAGE_EXT): return False
    return True

def product_image(p):
    try: imgs = p['images']
    except: return None
    if not imgs: return None
    if isinstance(imgs, str):
        try: imgs = json.loads(imgs)
        except: return None
    if isinstance(imgs, (list, tuple)):
        for u in imgs:
            if not isinstance(u, str): continue
            u_low = u.lower()
            if u_low.endswith(('.pdf', '.dwg', '.zip', '.doc', '.docx')): continue
            if any(x in u_low for x in ('схем', 'черт', 'drawing', 'schema')): continue
            if u.startswith('http') or u.startswith('/static/'): return u
    return None

app.jinja_env.globals['product_image'] = product_image

def _schema_kind(url):
    low = str(url or '').lower().split('?')[0].split('#')[0]
    if low.endswith('.pdf'): return 'pdf'
    if low.endswith('.dwg'): return 'dwg'
    for e in _IMG_EXT:
        if low.endswith(e): return 'img'
    return 'other'

def _schema_name(url, idx=0):
    try: path = unquote(str(url or '').split('?')[0].split('#')[0])
    except: path = str(url or '')
    name = path.rstrip('/').split('/')[-1] or f'файл-{idx + 1}'
    return (name[:57] + '...') if len(name) > 60 else name

app.jinja_env.globals['schema_kind'] = _schema_kind
app.jinja_env.globals['schema_name'] = _schema_name

def _normalize_category(name, raw_category):
    n, raw = (name or '').lower().strip(), (raw_category or '').lower().strip()
    haystack = n + ' || ' + raw

    if n:
        if n.startswith(('унитаз', 'чаша унитаза', 'биде', 'бачок')): return 'Унитазы и биде'
        if n.startswith(('сиденье', 'крышка')): return 'Сиденья для унитазов'
        if n.startswith(('трап', 'сифон')): return 'Трапы и сифоны'
        if n.startswith('ванна'): return 'Ванны'
        if n.startswith(('раковина', 'умывальник')): return 'Раковины'
        if n.startswith('мойка'): return 'Кухонные мойки'
        if n.startswith('смеситель'): return 'Смесители'
        if n.startswith('полотенцесушитель'): return 'Полотенцесушители'
        if n.startswith(('тумба', 'пенал', 'шкаф', 'зеркальный шкаф', 'комплект мебели')): return 'Мебель и столешницы'
        if n.startswith('зеркало'): return 'Зеркала'
        if n.startswith(('кнопка', 'панель смыва', 'инсталляция', 'система инсталляции')): return 'Инсталляции и кнопки'

    if 'полотенцесушител' in haystack: return 'Полотенцесушители'
    if 'инсталляц' in haystack: return 'Инсталляции и кнопки'
    if 'кнопк' in haystack and ('смыв' in haystack or 'унитаз' in haystack): return 'Инсталляции и кнопки'
    if 'люк' in haystack and 'инсталл' not in haystack: return 'Сантехнические люки'
    if any(x in haystack for x in ('трап', 'сифон', 'душевой канал', 'душевой лоток', 'душевой трап', 'канал водоотвод')): return 'Трапы и сифоны'
    if any(x in haystack for x in ('унитаз', 'биде', 'бачок', 'чаша унитаза')): return 'Сиденья для унитазов' if 'сидень' in haystack or 'крышк' in haystack else 'Унитазы и биде'
    if 'сидень' in haystack or 'крышк' in haystack: return 'Сиденья для унитазов'
    if 'ванна' in haystack or 'ванн ' in haystack or raw in ('ванны', 'ванна'): return 'Ванны'
    if 'раковин' in haystack or 'умывальник' in haystack: return 'Раковины'
    if 'мойка' in haystack or 'мойк' in haystack: return 'Кухонные мойки'
    if 'столешниц' in haystack: return 'Мебель и столешницы'
    if 'смесител' in haystack: return 'Смесители'
    if any(x in haystack for x in ('кабина', 'огражден', 'уголок', 'шторка', 'поддон', 'дверь в нишу', 'душевая дверь', 'душевые двери', 'душевые углы', 'душевые перегородки', 'перегородк', 'душевой угол', 'душевая шторка', 'боковая панель')): return 'Душевые ограждения'
    if any(x in haystack for x in ('душевая стойка', 'душевая система', 'душевая панель', 'душевой гарнитур', 'душевой комплект', 'душевой набор', 'гигиенический душ', 'душевая лейка', 'душевой шланг', 'душевая штанга', 'стойка душ', 'ручной душ', 'верхний душ', 'лейка гигиеническая')): return 'Душевые системы'
    if 'душ' in haystack: return 'Душевые системы'
    if 'зеркал' in haystack: return 'Зеркала'
    if any(x in haystack for x in ('тумба', 'шкаф', 'пенал', 'комплект мебели', 'комплекты мебели', 'мебель', 'мебел')): return 'Мебель и столешницы'
    if any(x in haystack for x in ('держател', 'мыльниц', 'крюч', 'ершик', 'дозатор', 'диспенсер', 'полка', 'стакан', 'набор аксессуаров', 'аксессуар', 'полотенц', 'смеситель гигиенический', 'шланг', 'излив', 'лоток')): return 'Аксессуары'

    if raw_category:
        rc = str(raw_category).strip()
        if rc.lower() in ('', '-', 'nan', 'none') or len(rc) > 50 or rc.lower().startswith(('sheet', 'лист')): return 'Прочее'
        return rc
    return 'Прочее'

def icon_for_category(category):
    if not category: return "🔧"
    c = str(category).lower()
    icons = {'унитаз': '🚽', 'биде': '🚽', 'инсталляц': '🚽', 'кнопк': '🚽', 'ванн': '🛁', 'раковин': '🪣', 'мойк': '🪣', 'столешн': '🪣', 'смесител': '🚰', 'кран': '🚰', 'излив': '🚰', 'душ': '🚿', 'лейк': '🚿', 'гигиен': '🚿', 'кабин': '🚿', 'огражден': '🚿', 'поддон': '🚿', 'угол': '🚿', 'шторк': '🚿', 'полотенц': '🔥', 'мебел': '🪞', 'тумб': '🪞', 'зеркал': '🪞', 'шкаф': '🪞', 'пенал': '🪞', 'аксессуар': '🧴', 'диспенсер': '🧴', 'мыльниц': '🧴', 'держател': '🧴', 'трап': '🌀', 'сифон': '🌀', 'выпуск': '🌀', 'сидень': '🪑', 'крышк': '🪑', 'люк': '🚪'}
    for kw, emoji in icons.items():
        if kw in c: return emoji
    return '🔧'

def get_db():
    db = getattr(g, '_database', None)
    if db is None:
        db = g._database = sqlite3.connect(DATABASE, timeout=30)
        db.row_factory = sqlite3.Row
    return db

@app.teardown_appcontext
def close_db(exc):
    db = getattr(g, '_database', None)
    if db is not None: db.close()

SCHEMA = """
CREATE TABLE IF NOT EXISTS users (id INTEGER PRIMARY KEY AUTOINCREMENT, email TEXT UNIQUE NOT NULL, password_hash TEXT NOT NULL, name TEXT DEFAULT '', phone TEXT DEFAULT '', address TEXT DEFAULT '', created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP);
CREATE TABLE IF NOT EXISTS products (id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT NOT NULL, category TEXT NOT NULL, brand TEXT NOT NULL, price REAL NOT NULL, old_price REAL, sku TEXT DEFAULT '', description TEXT DEFAULT '', icon TEXT DEFAULT '🔧', color TEXT DEFAULT 'from-blue-500 to-cyan-500', material TEXT DEFAULT '', diameter TEXT DEFAULT '', images TEXT DEFAULT '[]', schemas TEXT DEFAULT '[]', extra TEXT DEFAULT '{}', in_stock INTEGER DEFAULT 1, is_new INTEGER DEFAULT 0, is_hit INTEGER DEFAULT 0);
CREATE TABLE IF NOT EXISTS orders (id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER NOT NULL, total REAL NOT NULL, status TEXT DEFAULT 'Новый', address TEXT DEFAULT '', phone TEXT DEFAULT '', comment TEXT DEFAULT '', created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP, FOREIGN KEY (user_id) REFERENCES users (id));
CREATE TABLE IF NOT EXISTS order_items (id INTEGER PRIMARY KEY AUTOINCREMENT, order_id INTEGER NOT NULL, product_id INTEGER, product_name TEXT, quantity INTEGER NOT NULL, price REAL NOT NULL, FOREIGN KEY (order_id) REFERENCES orders (id));
CREATE TABLE IF NOT EXISTS favorites (user_id INTEGER NOT NULL, product_id INTEGER NOT NULL, PRIMARY KEY (user_id, product_id));
"""

DEMO_PRODUCT_NAMES = ['Смеситель для кухни Grohe Minta', 'Смеситель для ванны Hansgrohe Logis', 'Смеситель для раковины Grohe Eurosmart', 'Душевая система Grohe Rainshower', 'Душевой гарнитур Hansgrohe Croma', 'Гигиенический душ Iddis', 'Унитаз подвесной Roca Inspira', 'Унитаз-компакт Vitra S50', 'Раковина накладная Roca Uno', 'Раковина с тумбой Am.Pm', 'Ванна акриловая Roca Uno 170', 'Ванна чугунная 150 см', 'Полотенцесушитель водяной Лесенка', 'Полотенцесушитель электрический']

def init_db():
    db = get_db()
    db.executescript(SCHEMA)
    for col, default in (('images', "'[]'"), ('extra', "'{}'"), ('sku', "''"), ('schemas', "'[]'")):
        try: db.execute(f'ALTER TABLE products ADD COLUMN {col} TEXT DEFAULT {default}')
        except sqlite3.OperationalError: pass
    db.commit()

with app.app_context(): init_db()

def login_required(f):
    @wraps(f)
    def wrapper(*args, **kwargs):
        if 'user_id' not in session:
            flash('Войдите в аккаунт, чтобы продолжить', 'warning')
            return redirect(url_for('login', next=request.path))
        return f(*args, **kwargs)
    return wrapper

def get_cart(): return session.get('cart', {})

def save_cart(cart):
    session['cart'] = cart
    session.modified = True

def get_cart_items():
    cart = get_cart()
    if not cart: return [], 0
    db = get_db()
    items, total = [], 0
    for pid, qty in cart.items():
        p = db.execute('SELECT * FROM products WHERE id=?', (pid,)).fetchone()
        if p:
            sub = p['price'] * qty
            total += sub
            items.append({'product': p, 'quantity': qty, 'subtotal': sub})
    return items, total

def _product_to_dict(row):
    d = dict(row)
    for field in ('images', 'schemas'):
        try: d[field] = json.loads(d.get(field) or '[]')
        except: d[field] = []
    try: d['extra'] = json.loads(d.get('extra') or '{}')
    except: d['extra'] = {}
    return d

@app.context_processor
def inject_globals():
    user = None
    if 'user_id' in session:
        db = get_db()
        user = db.execute('SELECT * FROM users WHERE id=?', (session['user_id'],)).fetchone()
    db = get_db()
    ALLOWED_CATEGORIES = ['Унитазы и биде', 'Раковины', 'Инсталляции и кнопки', 'Полотенцесушители', 'Душевые ограждения', 'Трапы и сифоны', 'Смесители', 'Душевые системы', 'Мебель и столешницы', 'Душевые кабины', 'Ванны', 'Зеркала', 'Прочее']
    placeholders = ','.join('?' * len(ALLOWED_CATEGORIES))
    rows = db.execute(f'SELECT category FROM products WHERE category IN ({placeholders}) GROUP BY category', ALLOWED_CATEGORIES).fetchall()
    found_cats = {r['category'] for r in rows}
    cats = [c for c in ALLOWED_CATEGORIES if c in found_cats]
    brands = [r['brand'] for r in db.execute('SELECT brand, COUNT(*) AS cnt FROM products WHERE brand != "" GROUP BY brand ORDER BY cnt DESC').fetchall()]
    return {'cart_count': sum(get_cart().values()), 'current_user': user, 'categories': cats, 'brands_list': brands}

@app.route('/')
def index():
    db = get_db()
    hits = db.execute('SELECT * FROM products WHERE is_hit=1 LIMIT 8').fetchall()
    new = db.execute('SELECT * FROM products WHERE is_new=1 LIMIT 4').fetchall()
    return render_template('index.html', hits=hits, new=new)

def _is_white_image(url):
    if url in _WHITE_IMAGE_CACHE: return _WHITE_IMAGE_CACHE[url]
    try:
        r = requests.get(url, timeout=3, headers={'User-Agent': 'Mozilla/5.0'})
        if r.status_code != 200 or len(r.content) < 1000:
            _WHITE_IMAGE_CACHE[url] = True
            return True
        img = Image.open(io.BytesIO(r.content))
        if img.mode in ('RGBA', 'LA') or (img.mode == 'P' and 'transparency' in img.info):
            img = img.convert('RGBA')
            bg = Image.new('RGBA', img.size, (255, 255, 255, 255))
            bg.paste(img, mask=img.split()[3])
            img = bg.convert('RGB')
        else: img = img.convert('RGB')
        img = img.resize((32, 32))
        pixels = list(img.getdata())
        white_count = sum(1 for r_col, g_col, b_col in pixels if r_col > 235 and g_col > 235 and b_col > 235)
        is_white = (white_count / len(pixels)) > 0.60
        _WHITE_IMAGE_CACHE[url] = is_white
        return is_white
    except:
        _WHITE_IMAGE_CACHE[url] = True
        return True

@app.route('/collections')
def collections():
    db = get_db()
    allowed_categories = ['Душевые кабины и ограждения', 'Унитазы и биде', 'Раковины', 'Ванны', 'Смесители', 'Полотенцесушители', 'Мебель для ванной', 'Душевые системы', 'Зеркала', 'Аксессуары', 'Трапы и сифоны']
    cols = {}
    for cat_name in allowed_categories:
        prods = db.execute("SELECT images FROM products WHERE category=? AND images IS NOT NULL AND images != '[]' AND images != '' ORDER BY RANDOM() LIMIT 30", (cat_name,)).fetchall()
        images_list = []
        for p in prods:
            if len(images_list) >= 5: break
            try:
                imgs = json.loads(p['images'])
                if isinstance(imgs, list):
                    for img_url in imgs:
                        if not isinstance(img_url, str) or not img_url.startswith('http'): continue
                        if img_url.lower().endswith(('.pdf', '.dwg', '.zip', '.doc', '.docx')): continue
                        if _is_white_image(img_url): continue
                        if img_url not in images_list:
                            images_list.append(img_url)
                            if len(images_list) >= 5: break
            except: pass
        if images_list: cols[cat_name] = images_list
    return render_template('collections.html', collections=cols)

@app.route('/partners')
def partners(): return render_template('partners.html')

@app.route('/where-to-buy')
def where_to_buy(): return render_template('where_to_buy.html')

@app.route('/catalog')
def catalog():
    db = get_db()
    category = request.args.get('category', '')
    brand = request.args.get('brand', '')
    price_min = request.args.get('price_min', type=float)
    price_max = request.args.get('price_max', type=float)
    material = request.args.get('material', '')
    sort = request.args.get('sort', 'popular')
    q = request.args.get('q', '').strip()

    sql = 'SELECT * FROM products WHERE 1=1'
    params = []
    if category: sql += ' AND category=?'; params.append(category)
    if brand: sql += ' AND brand=?'; params.append(brand)
    if material: sql += ' AND material=?'; params.append(material)
    if price_min is not None: sql += ' AND price>=?'; params.append(price_min)
    if price_max is not None: sql += ' AND price<=?'; params.append(price_max)
    if q:
        sql += ' AND (name LIKE ? OR description LIKE ? OR sku LIKE ?)'
        params += [f'%{q}%', f'%{q}%', f'%{q}%']

    total_count = db.execute(sql.replace('SELECT *', 'SELECT COUNT(*)'), params).fetchone()[0]

    try: per_page = int(request.args.get('per_page', 24))
    except: per_page = 24
    if per_page not in (0, 24, 48, 96, 192): per_page = 24

    page = request.args.get('page', type=int, default=1)
    if page < 1: page = 1

    if per_page == 0:
        total_pages, page, start_p, end_p = 1, 1, 1, 1
    else:
        total_pages = max(1, (total_count + per_page - 1) // per_page)
        start_p, end_p = max(1, page - 2), min(total_pages, page + 2)

    if category == 'Унитазы и биде': order_sql = " ORDER BY CASE WHEN LOWER(brand) = 'irai' THEN 1 ELSE 2 END, is_hit DESC, id DESC"
    elif category in ('Душевые ограждения', 'Душевые кабины'): order_sql = " ORDER BY CASE WHEN LOWER(brand) = 'veconi' THEN 1 ELSE 2 END, is_hit DESC, id DESC"
    elif category == 'Раковины': order_sql = " ORDER BY CASE WHEN LOWER(brand) = 'ceramicanova' THEN 1 ELSE 2 END, is_hit DESC, id DESC"
    elif category == 'Мебель и столешницы': order_sql = " ORDER BY CASE WHEN LOWER(brand) = 'sancos' THEN 1 ELSE 2 END, is_hit DESC, id DESC"
    elif category == 'Смесители': order_sql = " ORDER BY CASE WHEN LOWER(brand) = 'aquame' THEN 1 ELSE 2 END, is_hit DESC, id DESC"
    else:
        order_sql = " ORDER BY CASE WHEN LOWER(brand) = 'irai' THEN 1 WHEN LOWER(brand) = 'ceramicanova' THEN 2 WHEN LOWER(brand) = 'veconi' THEN 3 WHEN LOWER(brand) = 'sancos' THEN 4 WHEN LOWER(brand) = 'aquame' THEN 5 ELSE 6 END, is_hit DESC, id DESC"
        if sort == 'price_asc': order_sql = " ORDER BY price ASC"
        elif sort == 'price_desc': order_sql = " ORDER BY price DESC"
        elif sort == 'new': order_sql = " ORDER BY is_new DESC, id DESC"

    sql += order_sql
    if per_page != 0: sql += f" LIMIT {per_page} OFFSET {(page - 1) * per_page}"

    products_page = db.execute(sql, params).fetchall()
    materials = [r['material'] for r in db.execute('SELECT DISTINCT material FROM products WHERE material != ""').fetchall()]

    qp = {k: v for k, v in [('q', q), ('category', category), ('brand', brand), ('material', material), ('price_min', price_min), ('price_max', price_max), ('sort', sort)] if v is not None and v != ''}
    base_query = urlencode(qp)
    query_prefix = ('?' + base_query + '&') if base_query else '?'

    return render_template('catalog.html', products=products_page, materials=materials, category=category, brand=brand, material=material, price_min=price_min, price_max=price_max, sort=sort, q=q, page=page, total_pages=total_pages, total_count=total_count, per_page=per_page, start_p=start_p, end_p=end_p, base_query=base_query, query_prefix=query_prefix)

def _get_base_name(name, color):
    if not name: return ""
    n = str(name).strip()
    parts = n.split(',')
    if len(parts) >= 2: return ",".join(parts[:2]).lower().strip()
    c = str(color or '').lower().strip()
    n_low = n.lower().replace(c, '') if c else n.lower()
    return re.sub(r'[\s,.\-]+', ' ', n_low).strip()

@app.route('/product/<int:pid>')
def product(pid):
    db = get_db()
    row = db.execute('SELECT * FROM products WHERE id=?', (pid,)).fetchone()
    if not row:
        flash('Товар не найден', 'error')
        return redirect(url_for('catalog'))
    p = _product_to_dict(row)
    
    variations = []
    base_n = _get_base_name(p['name'], p['color'])
    for bp in db.execute('SELECT * FROM products WHERE brand=?', (p['brand'],)).fetchall():
        if _get_base_name(bp['name'], bp['color']) == base_n: variations.append(_product_to_dict(bp))
            
    similar = db.execute('SELECT * FROM products WHERE category=? AND id!=? LIMIT 4', (p['category'], pid)).fetchall()
    in_fav = 'user_id' in session and db.execute('SELECT 1 FROM favorites WHERE user_id=? AND product_id=?', (session['user_id'], pid)).fetchone() is not None
    return render_template('product.html', p=p, similar=similar, in_fav=in_fav, variations=variations)

# --- ЧАСТЬ 1 ЗАКОНЧИЛАСЬ. СРАЗУ ПОСЛЕ ЭТОЙ СТРОКИ ВСТАВЬ ЧАСТЬ 2 ---
@app.route('/cart')
def cart():
    items, total = get_cart_items()
    return render_template('cart.html', items=items, total=total)

@app.route('/cart/pdf')
def cart_pdf():
    items, total = get_cart_items()
    if not items:
        flash('Корзина пуста', 'warning')
        return redirect(url_for('cart'))

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
    elements = []
    
    font_path = os.path.join(os.environ.get('WINDIR', 'C:\\Windows'), 'Fonts', 'arial.ttf')
    font_bold_path = os.path.join(os.environ.get('WINDIR', 'C:\\Windows'), 'Fonts', 'arialbd.ttf')
    try:
        pdfmetrics.registerFont(TTFont('Arial', font_path))
        pdfmetrics.registerFont(TTFont('Arial-Bold', font_bold_path))
        font_name, font_bold = 'Arial', 'Arial-Bold'
    except: font_name, font_bold = 'Helvetica', 'Helvetica-Bold'

    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name='Russian', fontName=font_name, fontSize=10, leading=14))
    styles.add(ParagraphStyle(name='RussianBold', fontName=font_bold, fontSize=16, spaceAfter=20))
    styles.add(ParagraphStyle(name='RussianSub', fontName=font_name, fontSize=10, textColor=colors.gray, spaceAfter=20))

    elements.append(Paragraph("Спецификация заказа", styles['RussianBold']))
    elements.append(Paragraph("Сайт NIKA Сантехника", styles['RussianSub']))

    data = [['Наименование', 'Артикул', 'Кол-во', 'Цена', 'Сумма']]
    for item in items:
        p = item['product']
        data.append([Paragraph(str(p['name']), styles['Russian']), Paragraph(str(p['sku'] or '—'), styles['Russian']), str(item['quantity']), f"{p['price']:.0f} ₽", f"{item['subtotal']:.0f} ₽"])
    data.append(['', '', '', 'ИТОГО:', f"{total:.0f} ₽"])

    t = Table(data, colWidths=[250, 80, 50, 70, 70])
    t.setStyle(TableStyle([('BACKGROUND', (0,0), (-1,0), colors.HexColor('#0c2472')), ('TEXTCOLOR', (0,0), (-1,0), colors.white), ('ALIGN', (0,0), (-1,-1), 'LEFT'), ('ALIGN', (2,0), (-1,-1), 'CENTER'), ('FONTNAME', (0,0), (-1,-1), font_name), ('FONTNAME', (0,0), (-1,0), font_bold), ('FONTNAME', (3,-1), (-1,-1), font_bold), ('BOTTOMPADDING', (0,0), (-1,0), 12), ('TOPPADDING', (0,0), (-1,0), 12), ('BACKGROUND', (0,1), (-1,-1), colors.HexColor('#f8fafc')), ('GRID', (0,0), (-1,-1), 1, colors.HexColor('#cbd5e1'))]))
    elements.append(t)
    doc.build(elements)
    buffer.seek(0)
    return send_file(buffer, as_attachment=True, download_name='specification.pdf', mimetype='application/pdf')

@app.route('/cart/send-order', methods=['POST'])
def send_order():
    manager_email = request.form.get('manager_email')
    if not manager_email: return jsonify({'ok': False, 'error': 'Выберите менеджера'})
    items, total = get_cart_items()
    if not items: return jsonify({'ok': False, 'error': 'Корзина пуста'})

    SMTP_SERVER, SMTP_PORT = "smtp.yandex.ru", 465
    SENDER_EMAIL, SENDER_PASSWORD = "твоя_почта@yandex.ru", "твой_пароль"

    html = f"<h2>Новый заказ</h2><p><b>Сумма:</b> {total} ₽</p><table border='1'><tr><th>Товар</th><th>Артикул</th><th>Кол-во</th></tr>"
    for item in items: html += f"<tr><td>{item['product']['name']}</td><td>{item['product']['sku']}</td><td>{item['quantity']}</td></tr>"
    html += "</table>"

    msg = MIMEMultipart()
    msg['From'], msg['To'], msg['Subject'] = SENDER_EMAIL, manager_email, f"Заказ на {total} ₽"
    msg.attach(MIMEText(html, 'html'))

    try:
        server = smtplib.SMTP_SSL(SMTP_SERVER, SMTP_PORT)
        server.login(SENDER_EMAIL, SENDER_PASSWORD)
        server.send_message(msg)
        server.quit()
        return jsonify({'ok': True, 'message': 'Заказ отправлен!'})
    except Exception as e:
        return jsonify({'ok': False, 'error': f'Ошибка: {str(e)}'})

@app.route('/cart/add/<int:pid>', methods=['POST'])
def cart_add(pid):
    qty = int(request.form.get('qty', 1))
    cart = get_cart()
    cart[str(pid)] = cart.get(str(pid), 0) + qty
    save_cart(cart)
    if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
        return jsonify({'ok': True, 'count': sum(cart.values())})
    flash('Товар добавлен', 'success')
    return redirect(request.referrer or url_for('cart'))

@app.route('/cart/update', methods=['POST'])
def cart_update():
    pid, qty = request.form.get('pid'), int(request.form.get('qty', 1))
    cart = get_cart()
    if qty <= 0: cart.pop(pid, None)
    else: cart[pid] = qty
    save_cart(cart)
    return redirect(url_for('cart'))

@app.route('/cart/remove/<pid>', methods=['POST'])
def cart_remove(pid):
    cart = get_cart()
    cart.pop(str(pid), None)
    save_cart(cart)
    return redirect(url_for('cart'))

@app.route('/checkout', methods=['GET', 'POST'])
@login_required
def checkout():
    items, total = get_cart_items()
    if not items:
        flash('Корзина пуста', 'warning')
        return redirect(url_for('catalog'))
    if request.method == 'POST':
        address, phone, comment = request.form.get('address', '').strip(), request.form.get('phone', '').strip(), request.form.get('comment', '').strip()
        if not address or not phone: flash('Заполните адрес и телефон', 'error')
        else:
            db = get_db()
            cur = db.execute('INSERT INTO orders (user_id, total, address, phone, comment) VALUES (?,?,?,?,?)', (session['user_id'], total, address, phone, comment))
            order_id = cur.lastrowid
            for it in items: db.execute('INSERT INTO order_items (order_id, product_id, product_name, quantity, price) VALUES (?,?,?,?,?)', (order_id, it['product']['id'], it['product']['name'], it['quantity'], it['product']['price']))
            db.commit()
            session['cart'] = {}
            flash(f'Заказ #{order_id} оформлен!', 'success')
            return redirect(url_for('account'))
    return render_template('checkout.html', items=items, total=total)

@app.route('/register', methods=['GET', 'POST'])
def register():
    if request.method == 'POST':
        email, password, name = request.form.get('email', '').strip().lower(), request.form.get('password', ''), request.form.get('name', '').strip()
        if not email or not password: flash('Заполните все поля', 'error')
        elif len(password) < 6: flash('Пароль минимум 6 символов', 'error')
        else:
            db = get_db()
            try:
                db.execute('INSERT INTO users (email, password_hash, name) VALUES (?,?,?)', (email, generate_password_hash(password), name))
                db.commit()
                user = db.execute('SELECT * FROM users WHERE email=?', (email,)).fetchone()
                session['user_id'] = user['id']
                flash(f'Добро пожаловать!', 'success')
                return redirect(url_for('index'))
            except sqlite3.IntegrityError: flash('Такой email уже есть', 'error')
    return render_template('register.html')

@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        email, password = request.form.get('email', '').strip().lower(), request.form.get('password', '')
        db = get_db()
        user = db.execute('SELECT * FROM users WHERE email=?', (email,)).fetchone()
        if user and check_password_hash(user['password_hash'], password):
            session['user_id'] = user['id']
            flash('Вы вошли', 'success')
            return redirect(request.args.get('next') or url_for('index'))
        flash('Неверный email или пароль', 'error')
    return render_template('login.html')

@app.route('/logout')
def logout():
    session.pop('user_id', None)
    flash('Вы вышли', 'info')
    return redirect(url_for('index'))

@app.route('/account')
@login_required
def account():
    db = get_db()
    orders = db.execute('SELECT * FROM orders WHERE user_id=? ORDER BY id DESC', (session['user_id'],)).fetchall()
    favs = db.execute('SELECT p.* FROM favorites f JOIN products p ON p.id=f.product_id WHERE f.user_id=?', (session['user_id'],)).fetchall()
    return render_template('account.html', orders=orders, favorites=favs)

@app.route('/account/profile', methods=['POST'])
@login_required
def profile_update():
    db = get_db()
    db.execute('UPDATE users SET name=?, phone=?, address=? WHERE id=?', (request.form.get('name', '').strip(), request.form.get('phone', '').strip(), request.form.get('address', '').strip(), session['user_id']))
    db.commit()
    flash('Профиль обновлён', 'success')
    return redirect(url_for('account'))

@app.route('/favorites/toggle/<int:pid>', methods=['POST'])
@login_required
def fav_toggle(pid):
    db = get_db()
    if db.execute('SELECT 1 FROM favorites WHERE user_id=? AND product_id=?', (session['user_id'], pid)).fetchone():
        db.execute('DELETE FROM favorites WHERE user_id=? AND product_id=?', (session['user_id'], pid))
        state = False
    else:
        db.execute('INSERT INTO favorites (user_id, product_id) VALUES (?,?)', (session['user_id'], pid))
        state = True
    db.commit()
    return jsonify({'ok': True, 'in_fav': state})

@app.route('/contacts')
def contacts(): return render_template('contacts.html')

@app.route('/admin/delete-product/<int:pid>', methods=['POST'])
def admin_delete_product(pid):
    db = get_db()
    db.execute('DELETE FROM products WHERE id = ?', (pid,))
    db.commit()
    flash('Товар удален', 'success')
    return redirect(request.referrer or url_for('catalog'))

@app.route('/admin/manage', methods=['GET', 'POST'])
def admin_manage():
    db = get_db()
    if request.method == 'POST':
        if request.form.get('action') == 'delete':
            ids = request.form.getlist('product_ids')
            if ids:
                db.execute(f'DELETE FROM products WHERE id IN ({",".join("?" * len(ids))})', ids)
                db.commit()
                flash(f'Удалено: {len(ids)}', 'success')
            else: flash('Ничего не выбрано', 'warning')
        return redirect(request.referrer or url_for('admin_manage'))

    q, no_price, no_photo = request.args.get('q', '').strip(), request.args.get('no_price'), request.args.get('no_photo')
    page, per_page = request.args.get('page', 1, type=int), 50

    sql, params = 'SELECT * FROM products WHERE 1=1', []
    if q: sql += ' AND (name LIKE ? OR sku LIKE ? OR brand LIKE ?)'; params += [f'%{q}%'] * 3
    if no_price: sql += ' AND (price <= 0 OR price IS NULL)'
    if no_photo: sql += " AND (images IS NULL OR images = '[]' OR images = '')"

    total_count = db.execute(sql.replace('SELECT *', 'SELECT COUNT(*)'), params).fetchone()[0]
    total_pages = max(1, (total_count + per_page - 1) // per_page)
    sql += f' ORDER BY id DESC LIMIT {per_page} OFFSET {(page - 1) * per_page}'
    
    return render_template('admin_manage.html', products=db.execute(sql, params).fetchall(), q=q, page=page, total_pages=total_pages, total_count=total_count, no_price=no_price, no_photo=no_photo)

@app.route('/admin/stats')
def admin_stats():
    db = get_db()
    total = db.execute('SELECT COUNT(*) FROM products').fetchone()[0]
    with_price = db.execute('SELECT COUNT(*) FROM products WHERE price > 0').fetchone()[0]
    with_images = db.execute("SELECT COUNT(*) FROM products WHERE images IS NOT NULL AND images != '[]' AND images != ''").fetchone()[0]
    with_schemas = db.execute("SELECT COUNT(*) FROM products WHERE schemas IS NOT NULL AND schemas != '[]' AND schemas != ''").fetchone()[0]
    placeholders = ','.join('?' * len(DEMO_PRODUCT_NAMES))
    demo_count = db.execute(f"SELECT COUNT(*) FROM products WHERE name IN ({placeholders})", DEMO_PRODUCT_NAMES).fetchone()[0]
    by_cat = db.execute("SELECT category, COUNT(*) AS total, SUM(CASE WHEN price > 0 THEN 1 ELSE 0 END) AS with_price, SUM(CASE WHEN images != '[]' AND images IS NOT NULL THEN 1 ELSE 0 END) AS with_img, SUM(CASE WHEN schemas != '[]' AND schemas IS NOT NULL THEN 1 ELSE 0 END) AS with_schemas FROM products GROUP BY category ORDER BY total DESC").fetchall()
    return render_template('admin_stats.html', total=total, with_price=with_price, without_price=total - with_price, with_images=with_images, without_images=total - with_images, with_schemas=with_schemas, demo_count=demo_count, by_cat=by_cat, no_price_sample=db.execute('SELECT id, name, category, brand, price, sku FROM products WHERE price <= 0 ORDER BY id LIMIT 50').fetchall(), no_img_sample=db.execute("SELECT id, name, category, brand, price, sku FROM products WHERE images IS NULL OR images = '[]' OR images = '' ORDER BY id LIMIT 50").fetchall())

@app.route('/admin/product/<int:pid>')
def admin_product(pid):
    db = get_db()
    row = db.execute('SELECT * FROM products WHERE id=?', (pid,)).fetchone()
    if not row: flash('Товар не найден', 'error'); return redirect(url_for('admin_stats'))
    return render_template('admin_product.html', p=_product_to_dict(row))

@app.route('/admin/clear-all', methods=['GET', 'POST'])
def admin_clear_all():
    db = get_db()
    total = db.execute('SELECT COUNT(*) FROM products').fetchone()[0]
    if request.method == 'POST':
        try:
            with open(os.path.join(BASE_DIR, 'products_backup_full.json'), 'w', encoding='utf-8') as bf:
                json.dump([dict(r) for r in db.execute('SELECT * FROM products').fetchall()], bf, ensure_ascii=False, indent=2)
        except: pass
        db.execute('DELETE FROM products'); db.commit()
        flash(f'База очищена. Удалено {total}.', 'success')
        return redirect(url_for('admin_stats'))
    return render_template('admin_clear.html', total=total)

@app.route('/admin/drop-demo', methods=['GET', 'POST'])
def admin_drop_demo():
    db = get_db()
    if request.method == 'POST':
        cur = db.execute(f"DELETE FROM products WHERE name IN ({','.join('?' * len(DEMO_PRODUCT_NAMES))})", DEMO_PRODUCT_NAMES)
        db.commit()
        flash(f'Удалено демо: {cur.rowcount}', 'success')
    return redirect(url_for('admin_stats'))

@app.route('/admin/import', methods=['GET', 'POST'])
def admin_import():
    if request.method == 'POST':
        files = [f for f in request.files.getlist('files') if f and f.filename]
        if not files: flash('Файлы не выбраны', 'error'); return redirect(url_for('admin_import'))
        db = get_db()
        if request.form.get('clear') == 'on':
            db.execute('DELETE FROM products'); db.commit()
            if os.path.isdir(IMG_DIR): shutil.rmtree(IMG_DIR, ignore_errors=True)
            os.makedirs(IMG_DIR, exist_ok=True)
        items_by_key, files_info, price_map, promo_map = _read_files_to_items(files, print_log=True)
        merged_products, merge_stats = _merge_all_items(items_by_key)
        total_added = 0
        for item in merged_products:
            try:
                db.execute('''INSERT INTO products (name, category, brand, price, old_price, sku, description, icon, color, material, diameter, images, schemas, extra, in_stock, is_new, is_hit) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,1,?,0)''', (item.get('name', ''), item.get('category', 'Прочее'), item.get('brand', 'Прочие'), item.get('price') or 0, item.get('old_price'), item.get('sku') or '', item.get('description') or '', item.get('icon') or '🔧', item.get('color') or 'from-blue-500 to-cyan-500', item.get('material') or '', item.get('diameter') or '', item.get('images') or '[]', item.get('schemas') or '[]', item.get('extra') or '{}', item.get('is_new') or 0))
                total_added += 1
            except: pass
        db.commit()
        for sku_low, pr in price_map.items():
            if sku_low: db.execute('UPDATE products SET price=? WHERE LOWER(sku)=? AND (price IS NULL OR price=0)', (pr, sku_low))
        for sku_low, pr in promo_map.items():
            if sku_low:
                base = price_map.get(sku_low)
                db.execute('UPDATE products SET price=?, old_price=? WHERE LOWER(sku)=?', (pr, base, sku_low)) if base else db.execute('UPDATE products SET old_price=price, price=? WHERE LOWER(sku)=?', (pr, sku_low))
        db.commit()
        flash(f'Загружено файлов: {len(files)}. Добавлено товаров: {total_added}', 'success')
        return redirect(url_for('admin_stats'))
    return render_template('admin_import.html', total=get_db().execute('SELECT COUNT(*) FROM products').fetchone()[0])

@app.route('/admin/apply-prices', methods=['GET', 'POST'])
def admin_apply_prices():
    if request.method == 'POST':
        files = [f for f in request.files.getlist('files') if f and f.filename]
        if not files: flash('Файлы не выбраны', 'error'); return redirect(url_for('admin_apply_prices'))
        price_map = {}
        for f in files:
            try:
                if f.filename.lower().endswith(".xlsx"): sheets = _read_xlsx_all_sheets(f.read(), 'price')
                elif f.filename.lower().endswith(".xls"): sheets = _read_xls_old(f.read(), 'price')
                else: continue
                for _, rows, _, _ in sheets:
                    for row in rows:
                        sku, pr = _extract_sku(row), _price_from_row(row)
                        if sku and pr is not None: price_map[str(sku).strip().lower()] = pr
            except: pass
        db = get_db()
        updated = 0
        for sku_low, pr in price_map.items():
            updated += db.execute('UPDATE products SET price=? WHERE LOWER(sku)=? AND (price IS NULL OR price=0)', (pr, sku_low)).rowcount
        db.commit()
        flash(f'Собрано {len(price_map)} цен. Обновлено: {updated}.', 'success')
        return redirect(url_for('admin_apply_prices'))
    db = get_db()
    return render_template('admin_apply_prices.html', without_price=db.execute('SELECT COUNT(*) FROM products WHERE price <= 0').fetchone()[0], with_price=db.execute('SELECT COUNT(*) FROM products WHERE price > 0').fetchone()[0])

# ─────────────────────────────────────────────────────────────
#  УТИЛИТЫ ПАРСЕРА (полный блок)
# ─────────────────────────────────────────────────────────────
def _clean(v):
    if v is None: return None
    s = str(v).strip()
    return None if not s or s.lower() in ("nan", "none", "null", "-") else s

def _to_float(v):
    if v is None or v == "": return None
    if isinstance(v, (int, float)): return float(v) if v else None
    s = str(v).strip().replace(" ", "").replace("\xa0", "").replace("₽", "").replace("руб", "").replace("р.", "")
    if any(x in s.lower() for x in ("запрос", "уточн", "нет", "n/a")): return None
    s = s.replace('.', '').replace(',', '.') if '.' in s and ',' in s else s.replace(',', '.')
    m = re.search(r"-?\d+(?:\.\d+)?", s)
    return float(m.group(0)) if m else None

def _norm_key(s): return re.sub(r'[\s_\-.,\xa0/"\'\n]+', '', str(s or "").lower())

def _find_col(row, *needles, exclude=()):
    norm_needles, norm_exclude = [_norm_key(n) for n in needles], [_norm_key(e) for e in exclude]
    for k, v in row.items():
        nk = _norm_key(k)
        if any(e in nk for e in norm_exclude): continue
        if any(n in nk for n in norm_needles):
            c = _clean(v)
            if c: return c
    return None

def _find_header_row(all_rows, max_scan=15):
    MAIN_KEYWORDS = ("номенклатура", "наименование", "название", "артикул", "кодтовара", "товарнаягруппа", "категория", "раздел", "ррц", "мрц", "цена", "стоимость", "бренд", "производитель", "ссылка", "фото", "изображение", "name", "price", "category", "product", "brand", "sku")
    best_idx, best_score = 0, -1
    for i, row in enumerate(all_rows[:max_scan]):
        if not row: continue
        non_empty = [v for v in row if v is not None and str(v).strip() != ""]
        if len(non_empty) < 3: continue
        score = len(non_empty)
        joined_nk = _norm_key(" ".join(str(v).lower() for v in non_empty))
        score += sum(15 for kw in MAIN_KEYWORDS if kw in joined_nk)
        if len(non_empty) >= 5: score += 5
        if len(non_empty) <= 2 and len(joined_nk) < 30: score -= 30
        if 'ррц' in joined_nk or 'мрц' in joined_nk: score += 10
        if score > best_score: best_score, best_idx = score, i
    return best_idx

def _combine_header(all_rows, header_idx):
    rows_to_merge, last_idx = [header_idx], header_idx
    if header_idx > 0 and sum(1 for v in all_rows[header_idx - 1] if v is not None and str(v).strip() != "") >= 3: rows_to_merge.insert(0, header_idx - 1)
    for step in range(1, 3):
        ni = header_idx + step
        if ni >= len(all_rows): break
        nxt = all_rows[ni]
        nonempty_idx = [j for j, v in enumerate(nxt) if v is not None and str(v).strip() != ""]
        if not nonempty_idx: break
        joined_nk = _norm_key(" ".join(str(v).lower() for v in nxt if v is not None and str(v).strip() != ""))
        has_key = any(k in joined_nk for k in ('ррц', 'мрц', 'вал', 'остаток', 'наскладах', 'впути', 'впуть', 'цена', 'стоимость', 'руб', 'опт', 'розница'))
        if has_key or (nonempty_idx[0] >= 2 and max((len(str(v)) for v in nxt if v is not None and str(v).strip() != ""), default=0) < 50 and 2 <= len(nonempty_idx) < len(nxt) * 0.6):
            rows_to_merge.append(ni); last_idx = ni
        else: break
    n_cols = max((len(all_rows[i]) for i in rows_to_merge if i < len(all_rows)), default=0)
    while n_cols > 0 and not any(i < len(all_rows) and n_cols - 1 < len(all_rows[i]) and all_rows[i][n_cols - 1] is not None and str(all_rows[i][n_cols - 1]).strip() != "" for i in rows_to_merge): n_cols -= 1
    headers = []
    for j in range(n_cols):
        parts = []
        for i in rows_to_merge:
            if i < len(all_rows) and j < len(all_rows[i]) and all_rows[i][j] is not None and str(all_rows[i][j]).strip() != "":
                s = str(all_rows[i][j]).strip()
                if not parts or parts[-1].lower() != s.lower(): parts.append(s)
        headers.append(" / ".join(parts).lower())
    return headers, last_idx

def _zip_headers(headers, values):
    result, seen = {}, {}
    for i, h in enumerate(headers):
        key = str(h or '').strip() or f"__col_{i}__"
        seen[key] = seen.get(key, 0) + 1
        result[f"{key}_{seen[key]}" if seen[key] > 1 else key] = values[i] if i < len(values) else None
    return result

def _safe_dump_extra(extra, limit=12000):
    if not isinstance(extra, dict): return '{}'
    try:
        s = json.dumps(extra, ensure_ascii=False)
        if len(s) <= limit: return s
    except: return '{}'
    prio_prefixes = ('url', 'ссылканаизображ', 'ссылканадиск', 'фото', 'изображ', 'image', 'photo', 'picture', 'каталог', 'catalog', 'схем', 'чертеж', 'инструкц', 'паспорт', 'яндекс', 'yandex', 'disk', 'pictogram', 'инфографик', 'видео', 'video', '3d', '__link__')
    def prio(k):
        kl = _norm_key(k)
        return next((i for i, p in enumerate(prio_prefixes) if p in kl), 99)
    kept = {}
    for k, v in sorted(extra.items(), key=lambda kv: (prio(kv[0]), str(kv[0]))):
        trial = dict(kept); trial[k] = v
        try:
            if len(json.dumps(trial, ensure_ascii=False)) <= limit: kept[k] = v
        except: pass
    return json.dumps(kept, ensure_ascii=False)

def _save_all_embedded_images(xlsx_bytes, file_tag):
    result = {}
    try: zf = zipfile.ZipFile(io.BytesIO(xlsx_bytes))
    except: return result
    NS = {'main': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main', 'rel': 'http://schemas.openxmlformats.org/package/2006/relationships', 'xdr': 'http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing', 'a': 'http://schemas.openxmlformats.org/drawingml/2006/main', 'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
    try: wb_root = ET.fromstring(zf.read('xl/workbook.xml')); wb_rels = ET.fromstring(zf.read('xl/_rels/workbook.xml.rels'))
    except: zf.close(); return result
    rid_to_idx = {sh.get(f'{{{NS["r"]}}}id'): idx for idx, sh in enumerate(wb_root.findall('.//main:sheets/main:sheet', NS)) if sh.get(f'{{{NS["r"]}}}id')}
    rid_to_file = {rel.get('Id'): rel.get('Target') for rel in wb_rels.findall('rel:Relationship', NS)}
    os.makedirs(IMG_DIR, exist_ok=True)
    for rid, sheet_idx in rid_to_idx.items():
        if not rid_to_file.get(rid): continue
        sheet_base = os.path.basename(rid_to_file[rid])
        try: sheet_rels = ET.fromstring(zf.read(f'xl/worksheets/_rels/{sheet_base}.rels'))
        except KeyError: continue
        drawing_target = next((rel.get('Target') for rel in sheet_rels.findall('rel:Relationship', NS) if 'drawing' in (rel.get('Type') or '').lower()), None)
        if not drawing_target: continue
        drawing_path = 'xl/' + drawing_target[3:] if drawing_target.startswith('../') else ('xl/' + drawing_target if not drawing_target.startswith('/') else drawing_target[1:])
        try: drawing_root = ET.fromstring(zf.read(drawing_path))
        except KeyError: continue
        drawing_base = os.path.basename(drawing_path)
        rid_to_media = {}
        try:
            draw_rels = ET.fromstring(zf.read(f'xl/drawings/_rels/{drawing_base}.rels'))
            rid_to_media = {rel.get('Id'): rel.get('Target') for rel in draw_rels.findall('rel:Relationship', NS) if 'image' in (rel.get('Type') or '').lower()}
        except KeyError: pass
        row_to_media = {}
        for anchor in drawing_root:
            if anchor.tag.split('}')[-1] not in ('oneCellAnchor', 'twoCellAnchor'): continue
            row_el = anchor.find('xdr:from/xdr:row', NS)
            blip = anchor.find('.//a:blip', NS)
            if row_el is not None and blip is not None and row_el.text is not None:
                media_target = rid_to_media.get(blip.get(f'{{{NS["r"]}}}embed'))
                if media_target:
                    media_path = 'xl/' + media_target[3:] if media_target.startswith('../') else ('xl/' + media_target if not media_target.startswith('/') else media_target[1:])
                    row_to_media.setdefault(int(row_el.text), []).append(media_path)
        if row_to_media:
            sheet_map = {}
            for row_idx, media_list in row_to_media.items():
                urls = []
                for i, media_path in enumerate(media_list):
                    try: data = zf.read(media_path)
                    except KeyError: continue
                    ext = os.path.splitext(media_path)[1].lower()
                    fname = re.sub(r'[^\w\-.]', '_', f"emb_{file_tag}_s{sheet_idx}_r{row_idx}_{i}{ext if ext in ('.png', '.jpg', '.jpeg', '.webp', '.gif') else '.png'}")
                    with open(os.path.join(IMG_DIR, fname), 'wb') as f: f.write(data)
                    urls.append(IMG_URL_PREFIX + fname)
                if urls: sheet_map[row_idx] = urls
            if sheet_map: result[sheet_idx] = sheet_map
    zf.close()
    return result

def _read_hyperlinks_from_xlsx_xml(xlsx_bytes):
    result = {}
    try: zf = zipfile.ZipFile(io.BytesIO(xlsx_bytes))
    except: return result
    NS = {'main': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main', 'rel': 'http://schemas.openxmlformats.org/package/2006/relationships', 'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
    try: wb_root = ET.fromstring(zf.read('xl/workbook.xml')); wb_rels = ET.fromstring(zf.read('xl/_rels/workbook.xml.rels'))
    except: zf.close(); return result
    rid_to_sheetfile = {rel.get('Id'): rel.get('Target') for rel in wb_rels.findall('rel:Relationship', NS)}
    for sheet_idx, sh in enumerate(wb_root.findall('.//main:sheets/main:sheet', NS)):
        target = rid_to_sheetfile.get(sh.get(f'{{{NS["r"]}}}id'))
        if not target: continue
        sheet_path = target[1:] if target.startswith('/') else 'xl/' + target
        try: sheet_root = ET.fromstring(zf.read(sheet_path))
        except: continue
        rid_to_url = {}
        try:
            sheet_rels = ET.fromstring(zf.read(f'xl/worksheets/_rels/{os.path.basename(sheet_path)}.rels'))
            rid_to_url = {rel.get('Id'): rel.get('Target') for rel in sheet_rels.findall('rel:Relationship', NS) if 'hyperlink' in (rel.get('Type') or '').lower()}
        except: pass
        hyperlinks_map = {}
        hl_root = sheet_root.find('main:hyperlinks', NS)
        if hl_root is not None:
            for hl in hl_root.findall('main:hyperlink', NS):
                url = rid_to_url.get(hl.get(f'{{{NS["r"]}}}id'))
                ref = hl.get('ref')
                if url and ref:
                    m = re.match(r'([A-Z]+)(\d+)', ref)
                    if m:
                        col_0based = 0
                        for ch in m.group(1): col_0based = col_0based * 26 + (ord(ch) - ord('A') + 1)
                        hyperlinks_map.setdefault(int(m.group(2)) - 1, {})[col_0based - 1] = url
        if hyperlinks_map: result[sheet_idx] = hyperlinks_map
    zf.close()
    return result

def _read_xlsx_all_sheets(file_bytes, file_tag):
    all_embedded = _save_all_embedded_images(file_bytes, file_tag)
    hyperlinks_by_sheet = _read_hyperlinks_from_xlsx_xml(file_bytes)
    wb = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    result = []
    for sheet_idx, ws in enumerate(wb.worksheets):
        all_rows_values = list(ws.iter_rows(values_only=True))
        if not all_rows_values: continue
        header_idx = _find_header_row(all_rows_values)
        headers, last_header_row = _combine_header(all_rows_values, header_idx)
        hyperlinks_by_row = hyperlinks_by_sheet.get(sheet_idx, {})
        images_by_row = all_embedded.get(sheet_idx, {})
        rows = []
        for r_idx, raw in enumerate(all_rows_values[last_header_row + 1:], start=last_header_row + 1):
            if all(v is None for v in raw): continue
            row_dict = _zip_headers(headers, list(raw[:len(headers)]) + [None] * (len(headers) - len(raw)))
            keys_list = list(row_dict.keys())
            row_dict['__links__'] = {keys_list[c0]: url for c0, url in hyperlinks_by_row.get(r_idx, {}).items() if c0 < len(keys_list)}
            rows.append(row_dict)
        result.append((ws.title, rows, images_by_row, last_header_row))
    wb.close()
    return result

def _read_xls_old(file_bytes, file_tag):
    try: book = xlrd.open_workbook(file_contents=file_bytes)
    except: return _read_xlsx_all_sheets(file_bytes, file_tag)
    result = []
    for sheet_idx in range(book.nsheets):
        sheet = book.sheet_by_index(sheet_idx)
        if sheet.nrows < 2: continue
        all_rows = [[sheet.cell_value(r, c) for c in range(sheet.ncols)] for r in range(sheet.nrows)]
        header_idx = _find_header_row(all_rows)
        headers, last_header_row = _combine_header(all_rows, header_idx)
        rows = []
        for raw in all_rows[last_header_row + 1:]:
            if all(v in ("", None) for v in raw): continue
            row = _zip_headers(headers, list(raw[:len(headers)]) + [""] * (len(headers) - len(raw)))
            row['__links__'] = {}
            rows.append(row)
        if rows: result.append((sheet.name, rows, {}, last_header_row))
    return result

_IMAGE_KEY_SOFT = ("изображ", "фото", "картинк", "image", "photo", "picture", "ссылканаизображ", "pictogram", "инфографик", "инфографика")
_IMAGE_KEY_HARD_EXCLUDE = ("видео", "video", "3d", "чертеж", "чертёж", "схем", "схема", "drawing", "инструкц", "паспорт", "руководств", "документ", "сертификат", "обложк", "cover", "родитель")

def _extract_images(row, links=None, resolve_yandex=False):
    urls, seen = [], set()
    def add(u):
        if not u: return
        u = str(u).strip()
        if not u or not u.startswith('http') or _url_looks_like_schema(u): return
        if _is_yandex_disk_url(u):
            if resolve_yandex:
                resolved = _yandex_disk_image(u)
                if not resolved: return
                u = resolved
        elif not u.lower().split('?')[0].split('#')[0].endswith(_IMG_EXT): return
        if not _is_composite_url(u) and u not in seen:
            seen.add(u); urls.append(u)
    def is_image_col(k):
        kl, kln = str(k).lower(), _norm_key(k)
        if any(bad in kl or bad in kln for bad in _IMAGE_KEY_HARD_EXCLUDE): return False
        return any(x in kl for x in _IMAGE_KEY_SOFT) or "каталог" in kl or "catalog" in kl or "яндексдиск" in kln or "ссылканадиск" in kln or re.search(r"url\s*_?\s*\d+", kln) or re.search(r"\burl\b", kln) or "morephoto" in kln or "детальнаякартинка" in kln
    def priority(k):
        kl, kln = str(k).lower(), _norm_key(k)
        if "pictogram" in kl or "pictogram" in kln or "инфографик" in kl or "инфографик" in kln: return (9, 0)
        m = re.search(r"ссылк[ае]?\s+на\s+(фото|изображение|картинку)\s*(\d+)", kl)
        if m: return (0, int(m.group(2)))
        m = re.search(r"url\s*[_\-]?\s*(\d+)", kln)
        if m: return (0, int(m.group(1)))
        m = re.search(r"(изображение\s*[_\-]?\s*вид|фото|photo|image|img|pic)\s*[_\-]?\s*(\d+)", kl)
        if m: return (0, int(m.group(2)))
        if "ссылканафото" in kln or "ссылканаизображ" in kln: return (0, 99)
        if "детальнаякартинка" in kln or "morephoto" in kln: return (1, 0)
        if "яндексдиск" in kln or "ссылканадиск" in kln: return (2, 0)
        if "изображ" in kl: return (3, 0)
        if "каталог" in kl or "catalog" in kl: return (5, 0)
        return (4, 0)
    for k, v in sorted([(k, v) for k, v in row.items() if is_image_col(k)], key=lambda kv: priority(kv[0])):
        if v is not None:
            for u in _split_into_urls(str(v)): add(u)
    if links:
        for k, u in sorted([(k, u) for k, u in links.items() if is_image_col(k)], key=lambda kv: priority(kv[0])):
            for pu in _split_into_urls(str(u)): add(pu)
    return urls

_SCHEMA_DOC_KEYWORDS = ("каталог", "catalog", "pdf", "dwg", "schema", "instruction", "manual", "drawing", "документ", "сертификат", "схем")
_SCHEMA_IMG_KEYWORDS = ("схем", "чертеж", "чертёж", "инструкц", "паспорт", "руководств", "technicaldrawing", "technicaldrawings", "3dobject")
_SCHEMA_EXCLUDE_KEYWORDS = ("видео", "video", "pictogram", "инфографик", "обложк", "cover")

def _extract_schemas(row, links=None):
    urls, seen = [], set()
    def add(u, allow_images=False):
        if not u: return
        u = str(u).strip().rstrip(';,.')
        if not u.startswith('http') or _is_composite_url(u): return
        ext = _url_ext(u.lower().split('?')[0].split('#')[0])
        if allow_images and ext not in (_IMG_EXT + _DOC_EXT): return
        if not allow_images and ext not in _DOC_EXT: return
        if u not in seen:
            seen.add(u); urls.append(u)
    def col_kind(k):
        kl, kln = str(k).lower(), _norm_key(k)
        if any(bad in kl or bad in kln for bad in _SCHEMA_EXCLUDE_KEYWORDS): return None
        if any(kw in kl or kw in kln for kw in _SCHEMA_IMG_KEYWORDS): return 'img'
        if any(kw in kl or kw in kln for kw in _SCHEMA_DOC_KEYWORDS): return 'doc'
        return None
    for k, v in row.items():
        if k == '__links__' or v is None: continue
        kind = col_kind(k)
        if kind:
            for u in _split_into_urls(str(v)): add(u, allow_images=(kind == 'img'))
        for u in _split_into_urls(str(v)):
            if _url_looks_like_schema(u): add(u, allow_images=True)
    if links:
        for k, u in links.items():
            if not u: continue
            if _url_looks_like_schema(u): add(u, allow_images=True)
            elif col_kind(k): add(u, allow_images=(col_kind(k) == 'img'))
    return urls

def _extract_sku(row):
    for k, v in row.items():
        if _norm_key(k) == 'артикул':
            c = _clean(v)
            if c: return c
    v = _find_col(row, "артикулпродукта", "кодтовара", "sku", exclude=("родитель", "parent", "зеркал", "смесител", "схем", "id", "значение", "сверять", "остатки", "кодтнвэд", "штрихкод"))
    return v or _find_col(row, "артикул", exclude=("значение", "сверять", "остатки", "кодтнвэд")) or _find_col(row, "id", exclude=("тнвэд", "штрих")) or ""

_EXTRA_SKIP_PREFIXES = ("назван", "наимен", "артикул", "кодтовара", "sku", "катег", "раздел", "подкатег", "бренд", "производ", "компани", "описан", "материал", "цвет", "иконка", "icon", "новинка", "цен", "стоим", "ррц", "мрц", "старая", "oldprice")

def _detect_brand_by_sku(sku):
    if not sku: return ""
    s = str(sku).upper().strip()
    if s.startswith(('AQM', 'AQUA', 'AM')): return 'AQUAme'
    if s.startswith('CN'): return 'Ceramicanova'
    if s.startswith(('E1-', 'VMO-', 'EVO-')): return 'Veconi'
    if s.startswith('AR'): return 'Sancos'
    if s.startswith(('RAV', 'RAVAL')): return 'Raval'
    if s.startswith(('RV', 'RVD')): return 'Veconi'
    if s.startswith(('FL', 'FLOR')): return 'Florentina'
    if s.startswith(('PTDW', 'PDW')): return 'Veconi'
    return ""

def _price_from_row(row):
    pr = _to_float(_find_col(row, "ррцпоакции", "мрцпоакции", "ценапоакции", "акционнаяцена", "ценаакции", "ценасоскидкой"))
    if pr and pr > 0: return pr
    pr = _to_float(_find_col(row, "ррц", "мрц", "минрозничная", "розничнаяцена", "ценаруб", "ценапродажи", "цена", exclude=("закуп", "опт", "old", "скидк", "поакции", "вал", "старая")))
    if pr and pr > 0: return pr
    pr = _to_float(_find_col(row, "стоимость", "price"))
    if pr and pr > 0: return pr
    for key in ('__col_3__', '__col_4__', '__col_5__', '__col_2__', '__col_6__', '__col_1__'):
        v = row.get(key)
        if v is None: continue
        sv = str(v).strip().lower()
        if sv in ('ррц', 'мрц', 'вал.', 'вал', 'вал ', 'на складах', 'остаток', 'в пути', 'товары', 'none', '', 'nan', 'руб', 'руб.', 'р.', 'руб ') or sv.startswith(('>', '<')) or len(sv) > 20: continue
        pr = _to_float(v)
        if pr and 10 <= pr <= 10_000_000: return pr
    return None

def _map_row(row, sheet_name="", embedded_images=None, links=None, series=""):
    name = _find_col(row, "номенклатура", "названиепродукта", "наименованиевпрайсе", "наименованиеэлемента", "полноенаименованиеизделия", "полноенаименование") or _find_col(row, "краткоенаименованиеизделия", "краткоенаименование") or _find_col(row, "короткоенаименование", "расширенноенаименование") or _find_col(row, "наименование", "название") or _find_col(row, "серия/коллекция", "серия", "коллекция", "модель") or _find_col(row, "name", "product")
    if not name: return None
    name_str = str(name).strip()
    if len(name_str) < 2: return None
    sku = str(_extract_sku(row) or "").strip()
    if not sku and len(name_str) < 20 and not re.search(r'\d', name_str) and not embedded_images and not _extract_schemas(row, links=links) and (_price_from_row(row) or 0) <= 0: return None
    category = _normalize_category(name_str, _find_col(row, "категория", "товарнаягруппа", "названиераздела", "раздел", "подкатегория", "тип", "category") or (sheet_name.strip() if sheet_name else ""))
    brand = _detect_brand_by_sku(sku) or _find_col(row, "бренд", "brand") or _find_col(row, "производитель") or _find_col(row, "компания") or "Прочие"
    price_promo = _to_float(_find_col(row, "ррцпоакции", "мрцпоакции", "ценапоакции", "ценасоскидкой", "акционнаяцена", "ценаакции"))
    price_base = _to_float(_find_col(row, "ррц", "мрц", "минрозничная", "розничнаяцена", "ценаруб", "ценапродажи", "цена", exclude=("закуп", "опт", "old", "скидк", "поакции", "вал"))) or _to_float(_find_col(row, "стоимость", "price")) or _price_from_row(row)
    price = price_promo if price_promo is not None else price_base
    old_price = price_base if price_promo is not None else _to_float(_find_col(row, "стараяцена", "oldprice", "ценадо", "ценабезскидки", exclude=("закуп", "опт")))
    color_from_name = _extract_color_from_name(name_str)
    color = _find_col(row, "цвет", "color") or color_from_name
    images = _extract_images(row, links=links)
    if embedded_images:
        images.extend(u for u in embedded_images if u not in images)
    extra = {k: c for k, v in row.items() if k != '__links__' and not any(_norm_key(k).startswith(p) for p in _EXTRA_SKIP_PREFIXES) for c in [_clean(v)] if c}
    if links: extra.update({f"__link__{k}": u for k, u in links.items()})
    if series: extra['серия'] = series
    if color_from_name: extra['цвет_названия'] = color_from_name
    return {"name": product_display_name({"name": name_str, "sku": sku}), "sku": sku, "category": category, "brand": brand, "price": price or 0, "old_price": old_price, "description": _find_col(row, "описание", "description") or "", "material": _find_col(row, "материал", "material") or "", "color": color, "icon": _find_col(row, "иконка", "icon") or icon_for_category(category), "is_new": 1 if str(_find_col(row, "статусновинка", "новинка", "new") or "").strip().lower() in ("да", "yes", "1", "true", "+", "новинка") else 0, "images": json.dumps(images, ensure_ascii=False), "schemas": json.dumps(_extract_schemas(row, links=links), ensure_ascii=False), "extra": _safe_dump_extra(extra, limit=12000)}

def _merge_products(items, sources):
    if not items: return None
    if len(items) == 1:
        merged = dict(items[0])
        try: extra = json.loads(merged.get('extra') or '{}')
        except: extra = {}
        extra['_sources'] = sources[:]
        merged['extra'] = _safe_dump_extra(extra, limit=12000)
        return merged
    indexed = sorted(zip(items, sources), key=lambda t: len(re.sub(r'\s+', ' ', str(t[0].get('name') or '')).strip()), reverse=True)
    base = dict(indexed[0][0])
    src_list = []
    for _, s in indexed:
        if s not in src_list: src_list.append(s)
    base['name'] = max((it.get('name') or '' for it, _ in indexed), key=lambda n: len(re.sub(r'\s+', ' ', str(n)).strip()))
    base['sku'] = next((it['sku'] for it, _ in indexed if it.get('sku')), base.get('sku'))
    if not base.get('category') or base['category'] == 'Прочее': base['category'] = next((it['category'] for it, _ in indexed if it.get('category') and it.get('category') != 'Прочее'), 'Прочее')
    if not base.get('brand') or base['brand'] == 'Прочие': base['brand'] = next((it['brand'] for it, _ in indexed if it.get('brand') and it.get('brand') != 'Прочие'), 'Прочие')
    best_price, best_old = 0, None
    for it, _ in indexed:
        p = it.get('price') or 0
        if p > best_price: best_price, best_old = p, it.get('old_price')
    base['price'], base['old_price'] = best_price, best_old
    base['description'] = max((it.get('description') or '' for it, _ in indexed), key=len)
    for field in ('material', 'color', 'diameter', 'icon'):
        cur = base.get(field) or ''
        if not cur or (field == 'color' and cur.startswith('from-')):
            base[field] = next((it.get(field) for it, _ in indexed if it.get(field) and not (field == 'color' and it.get(field).startswith('from-'))), cur)
    base['is_new'] = 1 if any(it.get('is_new') for it, _ in indexed) else 0
    def _merge_lists(field):
        out, seen = [], set()
        for it, _ in indexed:
            try: lst = json.loads(it.get(field) or '[]')
            except: lst = []
            for u in lst:
                if isinstance(u, str) and u not in seen:
                    seen.add(u); out.append(u)
        return out
    base['images'] = json.dumps(_merge_lists('images'), ensure_ascii=False)
    base['schemas'] = json.dumps(_merge_lists('schemas'), ensure_ascii=False)
    merged_extra = {}
    for it, _ in reversed(indexed):
        try: merged_extra.update(json.loads(it.get('extra') or '{}'))
        except: pass
    merged_extra['серия'] = next((json.loads(it.get('extra') or '{}').get('серия') for it, _ in indexed if json.loads(it.get('extra') or '{}').get('серия')), '')
    merged_extra['_sources'] = src_list
    base['extra'] = _safe_dump_extra(merged_extra, limit=12000)
    return base

def _read_files_to_items(files, print_log=True, limit_rows=None):
    all_sheets, files_info = [], []
    for file_idx, f in enumerate(files):
        fname, fname_low = f.filename, f.filename.lower()
        file_tag = re.sub(r'[^\w\-]', '_', os.path.splitext(fname)[0])[:40]
        try:
            if fname_low.endswith('.csv'): sheets_parsed = [('csv', [dict(r, __links__={}) for r in csv.DictReader(io.StringIO(f.read().decode('utf-8-sig', errors='ignore')))], {}, 0)]
            elif fname_low.endswith('.xlsx'): sheets_parsed = _read_xlsx_all_sheets(f.read(), file_tag)
            elif fname_low.endswith('.xls'): sheets_parsed = _read_xls_old(f.read(), file_tag)
            else: files_info.append({'name': fname, 'error': 'формат не поддержан', 'sheets': []}); continue
        except Exception as e: files_info.append({'name': fname, 'error': f'ошибка: {e}', 'sheets': []}); continue
        f_info = {'name': fname, 'error': None, 'sheets': []}
        for sn, rows, images_by_row, header_idx in sheets_parsed:
            all_sheets.append((file_idx, fname, sn, rows, images_by_row, header_idx))
            f_info['sheets'].append({'sheet_name': sn, 'total_rows': len(rows), 'header_row_idx': header_idx})
        files_info.append(f_info)
    price_map, promo_map = {}, {}
    for _, _, sn, rows, _, _ in all_sheets:
        if _is_price_sheet(sn):
            for row in rows:
                sku = _extract_sku(row)
                if not sku: continue
                sku_low = str(sku).strip().lower()
                promo = _to_float(_find_col(row, "ррцпоакции", "мрцпоакции", "ценапоакции", "ррцпоакции08", "ррцпоакциидо", "акционнаяцена"))
                base = _to_float(_find_col(row, "ррц", "мрц", exclude=("закуп", "опт", "старая", "old", "скидк", "поакции")))
                if promo is not None: promo_map[sku_low] = promo
                if base is not None: price_map[sku_low] = base
                if base is None and promo is not None: price_map[sku_low] = promo
    items_by_key = {}
    for k, (fi, fname, sn, rows, images_by_row, header_idx) in enumerate(all_sheets, 1):
        if _is_price_sheet(sn) or _is_service_sheet(sn) or not rows: continue
        if limit_rows and len(rows) > limit_rows: rows = rows[:limit_rows]
        valid_series = _detect_valid_series(rows)
        current_series = ""
        for idx, raw in enumerate(rows):
            emb = images_by_row.get(header_idx + 1 + idx) or []
            links = raw.get('__links__') or {}
            norm = {str(k).strip().lower(): v for k, v in raw.items() if k != '__links__'}
            if idx in valid_series: current_series = valid_series[idx]; continue
            row_name_guess = _find_col(norm, "номенклатура", "наименование", "название", "name") or ""
            if _is_section_marker(row_name_guess, norm): continue
            item = _map_row(norm, sn, embedded_images=emb, links=links, series=current_series)
            if not item: continue
            sku_low = (item.get('sku') or '').strip().lower()
            if sku_low:
                promo, base = promo_map.get(sku_low), price_map.get(sku_low)
                cur = item.get('price') or 0
                if promo is not None: item['price'], item['old_price'] = promo, base or cur or item.get('old_price')
                elif base is not None and not cur: item['price'] = base
            items_by_key.setdefault(sku_low or ('n:' + (item.get('name') or '').strip().lower()), []).append((item, fname))
    return items_by_key, files_info, price_map, promo_map

def _merge_all_items(items_by_key):
    merged_products, stats = [], {'total_groups': len(items_by_key), 'merged_from_2plus': 0, 'total_items_used': 0, 'with_images': 0, 'with_price': 0}
    for key, pairs in items_by_key.items():
        items, sources = [p[0] for p in pairs], [p[1] for p in pairs]
        merged = _merge_products(items, sources)
        if not merged: continue
        if len(items) > 1: stats['merged_from_2plus'] += 1
        stats['total_items_used'] += len(items)
        if json.loads(merged.get('images') or '[]'): stats['with_images'] += 1
        if (merged.get('price') or 0) > 0: stats['with_price'] += 1
        merged_products.append(merged)
    return merged_products, stats

_SKIP_SHEET_NAMES = ("шк", "штрих", "промо", "schet")
def _is_service_sheet(sheet_name): return any(x in str(sheet_name or "").lower().strip() for x in _SKIP_SHEET_NAMES)
def _is_price_sheet(sheet_name): return any(x in str(sheet_name or "").lower().strip() for x in ("ррц", "акци"))

_GENERIC_SERIES_WORDS = {'аксессуары', 'ванны', 'смесители', 'мебель', 'керамика', 'инсталляции', 'полотенцесушители', 'унитазы', 'раковины', 'системы', 'ограждения', 'двери', 'уголки', 'шторки', 'гарнитуры', 'лейки', 'трапы', 'сифоны', 'зеркала', 'сиденья', 'панели', 'запчасти', 'комплекты', 'кабины', 'поддоны', 'смесители', 'ванна', 'смеситель', 'полотенцесушитель', 'керамика'}

def _is_section_marker(name, row):
    if not name: return False
    n = str(name).strip()
    if not n or len(n) > 80: return False
    if _extract_sku(row): return False
    if (_price_from_row(row) or 0) > 0: return False
    n_low = n.lower()
    if re.match(r'^aqm[\s\u00a0]+[а-яёa-z]', n_low) or n_low in {'avola', 'lanzo', 'lorento', 'treviso', 'alba', 'orvieto', 'ragusa', 'trevi', 'cucina'}: return True
    if re.match(r'^\d+\s+продукция\s+', n_low) or re.match(r'^\d+\s+[A-ZА-ЯЁ]', n): return True
    if re.match(r'^(gr|hg|ax|is|kludi|rc)[\s\u00a0]+[а-яёa-z\d]', n_low): return True
    generic = {'товары', 'зеркала', 'раковины', 'душевые трапы', 'сливы переливы', 'встраиваемые системы', 'клавиши старые', 'комплекты 4 в 1', 'запчасти для мебели', 'напольные унитазы', 'подвесные унитазы', 'мебель', 'полотенцесушители', 'смесители', 'аксессуары', 'душевые системы', 'унитазы', 'ванны', 'керамика'}
    if n_low in generic: return True
    if len(n) < 50 and not re.search(r'\d', n) and any(kw in n_low for kw in ('системы', 'ограждения', 'двери', 'уголки', 'шторки', 'гарнитуры', 'лейки', 'инсталляции', 'керамика', 'мебель', 'смесители', 'полотенцесушители', 'аксессуары', 'ванны', 'унитазы', 'раковины', 'сиденья', 'зеркала', 'трапы', 'сифоны', 'клапаны', 'сливы', 'панели', 'запчасти', 'комплекты', 'кабины', 'сегменты', 'боковые', 'подвесные', 'напольные', 'арматура')): return True
    return False

def _extract_series_key(name):
    if not name: return ""
    n = str(name).strip()
    if re.match(r'^\d+\s+продукция', n.lower()): return ""
    m = re.match(r'^(AQM|GR|HG|AX|IS|KLUDI|RC)[\s\u00a0]+(.+)$', n, re.IGNORECASE)
    if m: return "" if m.group(2).strip().lower() in _GENERIC_SERIES_WORDS else m.group(2).strip()
    m = re.match(r'^\d+\s+[A-ZА-ЯЁ][a-zA-Zа-яё]*[\s\u00a0]+(.+)$', n)
    if m: return "" if m.group(1).strip().lower() in _GENERIC_SERIES_WORDS else m.group(1).strip()
    if re.match(r'^[A-ZА-ЯЁ][a-zа-яё]+$', n): return "" if n.lower() in _GENERIC_SERIES_WORDS else n
    return ""

def _series_keyword(series_name):
    if not series_name: return ""
    for w in re.split(r'[\s\u00a0]+', series_name.strip()):
        w_clean = re.sub(r'[^\w]', '', w, flags=re.UNICODE)
        if len(w_clean) >= 4: return w_clean[:4].lower()
    return ""

def _detect_valid_series(rows):
    markers = []
    for idx, raw in enumerate(rows):
        norm = {str(k).strip().lower(): v for k, v in raw.items() if k != '__links__'}
        name = _find_col(norm, "номенклатура", "наименование", "название", "name") or ""
        if name and _is_section_marker(name, norm):
            key = _extract_series_key(name)
            keyword = _series_keyword(key) if key else ""
            if keyword: markers.append((idx, str(name).strip(), keyword))
    valid = {}
    for i, (idx, name, keyword) in enumerate(markers):
        next_idx = markers[i + 1][0] if i + 1 < len(markers) else len(rows)
        for j in range(idx + 1, next_idx):
            norm = {str(k).strip().lower(): v for k, v in rows[j].items() if k != '__links__'}
            tname = _find_col(norm, "номенклатура", "наименование", "название", "name") or ""
            if tname and keyword in str(tname).lower(): valid[idx] = name; break
    return valid

_COLOR_VALUES = ('золото шлифованное', 'никель шлифованный', 'медь шлифованная', 'хром глянцевый', 'хром матовый', 'никель полированный', 'чёрный матовый', 'черный матовый', 'чёрный глянцевый', 'черный глянцевый', 'белый матовый', 'белый глянцевый', 'серый матовый', 'розовый матовый', 'чёрное золото', 'черное золото', 'розовое золото', 'золото матовое', 'золото полированное', 'brushed gold', 'brushed nickel', 'brushed copper', 'brushed bronze', 'brushed black chrome', 'matte black', 'matt black', 'matt white', 'glossy gold', 'polished gold', 'polished brass', 'gun metal', 'gunmetal', 'old brass', 'oil rubbed bronze', 'antique bronze', 'antique copper', 'antique silver', 'titanium black', 'silk black', 'оружейная сталь', 'брашированное золото', 'брашированный никель', 'брашированная бронза', 'античная бронза', 'античная медь', 'античное серебро', 'серо-бежевый матовый', 'красная бронза металлик', 'графит металлик', 'латунь металлик', 'хром', 'никель', 'сатиновый', 'сатин', 'металл', 'золото', 'медь', 'антик', 'белый', 'черный', 'чёрный', 'серый', 'steel', 'gold', 'chrome', 'nickel', 'bronze', 'copper', 'black', 'white', 'silver', 'beige', 'антрацит')

def _extract_color_from_name(name):
    if not name: return ""
    n = str(name).strip()
    if not n or ',' not in n: return ""
    candidate = n[n.rfind(',') + 1:].strip().rstrip('.,;:')
    if not candidate: return ""
    cand_low = candidate.lower()
    for c in _COLOR_VALUES:
        if cand_low == c or cand_low.startswith(c + ' ') or cand_low.startswith(c + '/'): return candidate
    if len(cand_low) <= 35 and not re.search(r'\d+\s*(мм|mm|х|×|\*|x|X)', cand_low) and not re.search(r'\d{3,}', cand_low) and not any(s in cand_low for s in ('отв', 'смес', 'компл', 'внеш', 'встр', 'настен', 'напольн')):
        return candidate
    return ""

if __name__ == '__main__':
    port = int(os.environ.get('PORT', 5000))
    os.makedirs(IMG_DIR, exist_ok=True)
    app.run(host='0.0.0.0', port=port, debug=False)